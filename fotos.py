import csv
import io
import json
import logging
from functools import lru_cache
import os
from pathlib import Path
import re
import time
from urllib.parse import parse_qs, urlsplit
import urllib.request

from openpyxl import load_workbook

FOTOS_ARQUIVO = Path(__file__).resolve().parent / 'dados' / 'Fotos faltantes - Catálogo Munguba.xlsx'
FOTOS_COMPLEMENTARES = FOTOS_ARQUIVO.parent / 'fotos_complementares.json'
FOTOS_MINIATURAS = FOTOS_ARQUIVO.parent / 'fotos_miniaturas.json'

# Planilha pública com as fotos (exportada como CSV pelo próprio Google Sheets)
FOTOS_PLANILHA_ID = os.environ.get("FOTOS_PLANILHA_ID", "1nvX6uISakRKveiDBx6kjxi1OcRsw9BcXY4uzqHEoIR8")
FOTOS_PLANILHA_GID = os.environ.get("FOTOS_PLANILHA_GID", "1140076266")
FOTOS_CSV_URL = f"https://docs.google.com/spreadsheets/d/{FOTOS_PLANILHA_ID}/export?format=csv&gid={FOTOS_PLANILHA_GID}"
CACHE_SEGUNDOS = 60 * 60  # recarrega a planilha no máximo 1x por hora

# Epítetos que não identificam uma espécie (ex: "Inga sp. 1") -> não usar o binômio para casar
EPITETOS_INDEFINIDOS = {"sp", "sp.", "spp", "spp.", "cf", "cf.", "aff", "aff."}

_cache = {"fotos": None, "binomios": None, "carregado_em": 0}
logger = logging.getLogger(__name__)
_manifestos = {}


@lru_cache(maxsize=4)
def _ler_manifesto(caminho, modificado_em):
    return json.loads(Path(caminho).read_text(encoding='utf-8'))


def _manifesto(caminho):
    agora = time.monotonic()
    recente = _manifestos.get(caminho)
    if recente and agora - recente[0] < 1:
        return recente[1]
    try:
        modificado_em = caminho.stat().st_mtime_ns
    except FileNotFoundError:
        dados = {}
    else:
        dados = _ler_manifesto(str(caminho), modificado_em)
    _manifestos[caminho] = (agora, dados)
    return dados


def otimizar_foto(foto):
    """Usa versões pré-geradas; nunca baixa ou processa imagens na requisição."""
    local = _manifesto(FOTOS_MINIATURAS).get(foto['url'])
    return {**foto, **local} if local else foto


def normalizar_nome(nome):
    return re.sub(r"\s+", " ", str(nome or "")).strip().lower()


def binomio(nome):
    """Gênero + epíteto, sem autoria. Retorna None se o nome não tiver um epíteto válido."""
    partes = normalizar_nome(nome).split(" ")
    if (len(partes) < 2 or partes[1] in EPITETOS_INDEFINIDOS
            or not re.fullmatch(r"[^\W\d_]+(?:-[^\W\d_]+)*", partes[1])):
        return None
    return f"{partes[0]} {partes[1]}"


def _link_de_imagem(url):
    # Links .dzi (Deep Zoom) não são imagens que o navegador consegue exibir
    url = (url or "").strip()
    partes = urlsplit(url)
    if partes.scheme not in ('http', 'https') or not partes.netloc:
        return ""
    # O serviço FSI entrega uma imagem raster mesmo quando a origem é um DZI.
    raster_fsi = partes.path.endswith('/fsi/server') and parse_qs(partes.query).get('type') == ['image']
    if ".dzi" in url.lower() and not raster_fsi:
        return ""
    return url.replace("http://", "https://", 1)


def _carregar_planilha():
    with urllib.request.urlopen(FOTOS_CSV_URL, timeout=10) as resp:
        texto = resp.read().decode("utf-8-sig")

    leitor = csv.DictReader(io.StringIO(texto))
    if not leitor.fieldnames or "Nome científico (catálogo)" not in leitor.fieldnames:
        raise ValueError("A aba de fotos não contém o cabeçalho esperado")
    return _indexar_linhas(leitor)


def _carregar_arquivo():
    """Lê a cópia fornecida pelo usuário, sem modificar a planilha."""
    workbook = load_workbook(FOTOS_ARQUIVO, read_only=True, data_only=True)
    try:
        for aba in workbook:
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas, ())
            if "Nome científico (catálogo)" in cabecalho:
                return _indexar_linhas(dict(zip(cabecalho, valores)) for valores in linhas)
        raise ValueError("A planilha de fotos não contém o cabeçalho esperado")
    finally:
        workbook.close()


def _indexar_linhas(linhas):
    fotos, binomios = {}, {}
    for linha in linhas:
        nome = linha.get("Nome científico (catálogo)") or ""
        url_alta = _link_de_imagem(linha.get("Link da foto (alta resolução)"))
        url_media = _link_de_imagem(linha.get("Link da foto (tamanho médio)"))
        if not nome.strip() or not (url_alta or url_media):
            continue

        foto = {
            "url": url_alta or url_media,
            "url_media": url_media or url_alta,
            "autor": (linha.get("Autor/direitos") or "").strip(),
            "licenca": (linha.get("Licença") or "").strip(),
            "fonte": (linha.get("Página de referência") or "").strip(),
        }
        fotos.setdefault(normalizar_nome(nome), foto)
        chave_binomio = binomio(nome)
        if chave_binomio:
            binomios.setdefault(chave_binomio, foto)
    return fotos, binomios


def _garantir_cache():
    agora = time.time()
    if _cache["fotos"] is not None and agora - _cache["carregado_em"] < CACHE_SEGUNDOS:
        return
    try:
        carregar = _carregar_arquivo if FOTOS_ARQUIVO.is_file() else _carregar_planilha
        _cache["fotos"], _cache["binomios"] = carregar()
        logger.info("%s fotos carregadas da planilha.", len(_cache['fotos']))
    except Exception as e:
        # Se o Google estiver fora, mantém o cache antigo (ou vazio) e o site segue com a foto padrão
        logger.warning("Não foi possível carregar a planilha de fotos: %s", e)
        if _cache["fotos"] is None:
            _cache["fotos"], _cache["binomios"] = {}, {}
        # Não deixar uma falha temporária ocultar todas as fotos por uma hora.
        _cache["carregado_em"] = agora - CACHE_SEGUNDOS + 60
        return
    _cache["carregado_em"] = agora


def buscar_foto(nome_cientifico):
    """Procura a foto pelo nome científico exato; se não achar, tenta pelo binômio (ignora a autoria)."""
    _garantir_cache()
    foto = _manifesto(FOTOS_COMPLEMENTARES).get(normalizar_nome(nome_cientifico))
    if not foto:
        foto = _cache["fotos"].get(normalizar_nome(nome_cientifico))
    if not foto:
        chave_binomio = binomio(nome_cientifico)
        foto = _cache["binomios"].get(chave_binomio) if chave_binomio else None
    return otimizar_foto(foto) if foto else None
