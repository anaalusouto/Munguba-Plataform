"""Leitura dos Appendix S1, sem alterar as planilhas ou o banco legado."""
import json
import math
from collections import Counter
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

DADOS_DIR = Path(__file__).resolve().parent / 'dados'


def caminho_appendix(idioma):
    if idioma not in ('pt', 'en'):
        raise ValueError('Idioma inválido')
    return DADOS_DIR / f'Appendix_S1_Occurrence_records_{idioma.upper()}.xlsx'


def coordenadas_validas(latitude, longitude):
    if latitude is None or longitude is None:
        return False
    try:
        lat, lng = float(latitude), float(longitude)
    except (TypeError, ValueError):
        return False
    return math.isfinite(lat) and math.isfinite(lng) and -90 <= lat <= 90 and -180 <= lng <= 180


def texto(valor):
    return '' if valor is None else str(valor).strip()


@lru_cache(maxsize=4)
def _ler_appendix(caminho, modificado_em):
    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = workbook['Occurrences'].iter_rows(values_only=True)
        cabecalho = next(linhas)
        obrigatorias = {'occurrenceID', 'scientificName', 'acceptedNameUsage', 'dynamicProperties',
                        'decimalLatitude', 'decimalLongitude'}
        if not obrigatorias.issubset(cabecalho):
            raise ValueError('Appendix sem os campos Darwin Core obrigatórios')
        registros, ids = [], set()
        for numero, valores in enumerate(linhas, 2):
            if not any(v is not None for v in valores):
                continue
            registro = dict(zip(cabecalho, valores))
            identificador = texto(registro['occurrenceID'])
            if not identificador or identificador in ids:
                raise ValueError(f'occurrenceID ausente ou duplicado na linha {numero}')
            if not texto(registro['scientificName']):
                raise ValueError(f'scientificName ausente na linha {numero}')
            ids.add(identificador)
            registro['propriedades'] = json.loads(registro['dynamicProperties'] or '{}')
            registros.append(registro)
        vocabularios = {'bioeconomicUse': [], 'bioeconomicPlantPart': []}
        for chave, termo, equivalente, *_ in workbook['Vocabularies'].iter_rows(min_row=2, values_only=True):
            if chave in vocabularios:
                vocabularios[chave].append((texto(termo), texto(equivalente)))
        return registros, vocabularios
    finally:
        workbook.close()


def carregar_appendix(idioma):
    caminho = caminho_appendix(idioma)
    registros, vocabularios = _ler_appendix(str(caminho), caminho.stat().st_mtime_ns)
    # As duas abas Vocabularies apresentam PT na segunda coluna e EN na terceira.
    coluna = 1 if idioma == 'en' else 0
    filtros = {
        'categorias': [termos[coluna] for termos in vocabularios['bioeconomicUse']],
        'partes': [termos[coluna] for termos in vocabularios['bioeconomicPlantPart']],
    }
    canonicos = {
        'tags': {termo.casefold(): termo for termo in filtros['categorias']},
        'partes_tags': {termo.casefold(): termo for termo in filtros['partes']},
    }
    grupos = {}
    for registro in registros:
        nome = texto(registro['acceptedNameUsage'] or registro['scientificName'])
        chave = nome.casefold()
        planta = grupos.setdefault(chave, {
            'nome_cientifico': nome, 'familia': texto(registro.get('family')) or '-',
            'nomes_populares': [], 'tags': [], 'partes_tags': [], 'habitats': [],
            'formas_vida': [], 'referencias': [], 'pontos_gps': [], 'ocorrencias': [],
        })
        props = registro['propriedades']
        for campo, valores in (
            ('nomes_populares', texto(registro.get('vernacularName')).split(';')),
            ('tags', props.get('bioeconomicUse', [])),
            ('partes_tags', props.get('bioeconomicPlantPart', [])),
            ('formas_vida', props.get('lifeForm', [])),
            ('referencias', props.get('bioeconomicReferences', [])),
            ('habitats', [registro.get('habitat')]),
        ):
            for valor in valores:
                valor = texto(valor)
                valor = canonicos.get(campo, {}).get(valor.casefold(), valor)
                if valor and valor not in planta[campo]:
                    planta[campo].append(valor)
        nivel = props.get('georeferenceLevel', 'unknown')
        ponto = None
        if coordenadas_validas(registro['decimalLatitude'], registro['decimalLongitude']):
            ponto = {
                'lat': float(registro['decimalLatitude']), 'lng': float(registro['decimalLongitude']),
                'id': texto(registro['occurrenceID']), 'nivel': nivel,
                'incerteza': texto(registro.get('coordinateUncertaintyInMeters')),
            }
            planta['pontos_gps'].append(ponto)
        planta['ocorrencias'].append({
            'id': texto(registro['occurrenceID']), 'nome_original': texto(registro['scientificName']),
            'fonte': texto(registro.get('datasetName')), 'data': texto(registro.get('eventDate')),
            'tipo': texto(registro.get('basisOfRecord')), 'nivel': nivel, 'ponto': ponto,
            'localidade': texto(registro.get('locality')), 'habitat': texto(registro.get('habitat')),
            'referencia': texto(registro.get('bibliographicCitation')),
            'observacao_geo': texto(registro.get('georeferenceRemarks')),
            'verificacao': texto(registro.get('georeferenceVerificationStatus')),
        })
    plantas = sorted(grupos.values(), key=lambda p: p['nome_cientifico'].casefold())
    categorias = Counter(tag for p in plantas for tag in p['tags'])
    niveis = Counter(r['propriedades'].get('georeferenceLevel') for r in registros)
    stats = {
        'total_especies': len(plantas), 'total_ocorrencias': len(registros),
        'total_geo': sum(bool(p['pontos_gps']) for p in plantas),
        'total_pontos': sum(len(p['pontos_gps']) for p in plantas),
        'total_ocorrencia_geo': niveis['occurrence'], 'total_area_geo': niveis['area'],
        'por_categoria': dict(categorias),
    }
    return plantas, stats, filtros
